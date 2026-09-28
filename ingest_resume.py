"""Ingest a user-provided resume and prepare it for embedding and job matching."""

from __future__ import annotations

import argparse
import html
import json
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from dotenv import load_dotenv


OUTPUT_PATH = Path("cleaned_resume.json")


@dataclass(frozen=True)
class ResumeDocument:
    source_path: str
    file_name: str
    file_type: str
    raw_text: str
    cleaned_text: str
    embedding_input: str


def _normalize_text(value: str) -> str:
    text = unicodedata.normalize("NFKC", html.unescape(value))
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[\t\f\v]+", " ", text)
    text = re.sub(r"[ ]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r" *\n *", "\n", text)
    return text.strip()


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "PDF resume support requires pypdf. Install dependencies from requirements.txt."
        ) from exc

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "DOCX resume support requires python-docx. Install dependencies from requirements.txt."
        ) from exc

    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    return "\n".join(paragraphs)


def extract_resume_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".rtf"}:
        return _read_text_file(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)

    raise ValueError(f"Unsupported resume file type: {suffix or 'unknown'}")


def build_resume_document(path: Path) -> ResumeDocument:
    raw_text = extract_resume_text(path)
    cleaned_text = _normalize_text(raw_text)
    embedding_input = cleaned_text
    return ResumeDocument(
        source_path=str(path.resolve()),
        file_name=path.name,
        file_type=path.suffix.lower().lstrip("."),
        raw_text=raw_text,
        cleaned_text=cleaned_text,
        embedding_input=embedding_input,
    )


def save_resume_document(document: ResumeDocument, output_path: Path = OUTPUT_PATH) -> Path:
    payload: dict[str, Any] = {
        "source_path": document.source_path,
        "file_name": document.file_name,
        "file_type": document.file_type,
        "cleaned_text": document.cleaned_text,
        "embedding_input": document.embedding_input,
    }
    output_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Ingest a resume file for downstream embedding.")
    parser.add_argument("resume_path", help="Path to the resume file (.txt, .md, .pdf, or .docx).")
    parser.add_argument(
        "--output",
        default=str(OUTPUT_PATH),
        help="Path to write the cleaned resume JSON file.",
    )
    return parser.parse_args()


def main() -> int:
    load_dotenv()
    args = parse_args()
    resume_path = Path(args.resume_path)
    output_path = Path(args.output)

    if not resume_path.exists():
        print(f"Resume file not found: {resume_path}")
        return 1

    try:
        document = build_resume_document(resume_path)
        save_resume_document(document, output_path)
    except (ValueError, RuntimeError) as exc:
        print(str(exc))
        return 1

    print(f"Saved cleaned resume to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())